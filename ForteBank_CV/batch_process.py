r"""Batch extraction of annotated evidence frames for the ForteBank CV module.

Run from ForteBank_CV:
    .\.venv\Scripts\python.exe batch_process.py --excel "ForteBank_ML_timecodes_MMSS(1).xlsx"

Requires: opencv-python, openpyxl and the existing evidence_generator.py,
video_processor.py files. Local data and generated evidence should not be committed.
"""
import argparse
import csv
import json
import re
from datetime import datetime, time, timedelta
from pathlib import Path

import openpyxl
from evidence_generator import evidence_for_event


def time_to_seconds(value):
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, timedelta):
        return value.total_seconds()
    if isinstance(value, (datetime, time)):
        return value.hour * 3600 + value.minute * 60 + value.second + value.microsecond / 1_000_000
    if isinstance(value, (int, float)):
        # Excel time values can be fractions of a day.
        if 0 < value < 1:
            return round(float(value) * 86400, 3)
        return float(value)
    raw = str(value).strip().replace(",", ".")
    parts = raw.split(":")
    if len(parts) == 2:
        return int(parts[0]) * 60 + float(parts[1])
    if len(parts) == 3:
        return int(parts[0]) * 3600 + int(parts[1]) * 60 + float(parts[2])
    return float(raw)


def video_key(value):
    name = Path(str(value).strip()).stem if str(value).lower().endswith(('.mov', '.mp4', '.avi')) else str(value).strip()
    name = name.upper().replace('Х', 'X')
    match = re.fullmatch(r'X\s*-\s*0*(\d+)', name)
    return f'X-{int(match.group(1))}' if match else name


def main():
    parser = argparse.ArgumentParser(description='Extract annotated ForteBank video evidence')
    parser.add_argument('--excel', default='ForteBank_ML_timecodes_MMSS(1).xlsx')
    parser.add_argument('--videos', default='videos')
    parser.add_argument('--output', default='evidence')
    args = parser.parse_args()

    excel_path = Path(args.excel)
    video_dir = Path(args.videos)
    output_dir = Path(args.output)
    if not excel_path.is_file():
        parser.error(f'Excel file not found: {excel_path}')
    if not video_dir.is_dir():
        parser.error(f'Video directory not found: {video_dir}')

    wb = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
    try:
        if 'AI_Annotations' not in wb.sheetnames:
            parser.error('Worksheet AI_Annotations not found')
        sheet = wb['AI_Annotations']
        rows = sheet.iter_rows(values_only=True)
        headers = [str(h).strip().lower() if h is not None else '' for h in next(rows)]
        required = ['video_id', 'timecode (mm:ss)', 'verdict', 'violation_id', 'violation_type']
        missing = [c for c in required if c not in headers]
        if missing:
            parser.error(f'Missing Excel columns: {missing}')
        idx = {name: headers.index(name) for name in required}
        description_idx = headers.index('original_description') if 'original_description' in headers else None
        annotations = list(rows)
    finally:
        wb.close()

    files = sorted((p for p in video_dir.iterdir() if p.is_file() and p.suffix.lower() == '.mov'))
    by_key = {}
    for file in files:
        by_key.setdefault(video_key(file.stem), []).append(file)

    output_dir.mkdir(parents=True, exist_ok=True)
    report = []
    results = []
    counts = {'ok': 0, 'skipped_no_violation': 0, 'missing_video': 0, 'missing_timecode': 0, 'ambiguous_video': 0, 'error': 0}

    for excel_row, row in enumerate(annotations, start=2):
        vid = row[idx['video_id']]
        if vid is None:
            continue
        key = video_key(vid)
        verdict = str(row[idx['verdict']] or '').strip().lower()
        record = {'excel_row': excel_row, 'video_id': key, 'verdict': verdict}
        if verdict != 'violation':
            status = 'skipped_no_violation'
        else:
            try:
                seconds = time_to_seconds(row[idx['timecode (mm:ss)']])
                if seconds is None:
                    status = 'missing_timecode'
                elif key not in by_key:
                    status = 'missing_video'
                elif len(by_key[key]) != 1:
                    status = 'ambiguous_video'
                else:
                    video_path = by_key[key][0]
                    # evidence_for_event writes JPEGs to the current working directory's
                    # evidence/ folder. Run from ForteBank_CV; outputs are unique by video stem/time.
                    frames = evidence_for_event(video_path, seconds)
                    results.append({
                        'video_id': key,
                        'video_filename': video_path.name,
                        'verdict': 'violation',
                        'violation_id': row[idx['violation_id']],
                        'violation_type': row[idx['violation_type']],
                        'description': row[description_idx] if description_idx is not None else None,
                        'timecode_sec': seconds,
                        'evidence_frames': frames,
                    })
                    status = 'ok'
            except Exception as exc:
                status = 'error'
                record['error'] = str(exc)
        counts[status] += 1
        record['status'] = status
        report.append(record)
        print(f"{key}: {status}" + (f" ({record['error']})" if 'error' in record else ''))

    (output_dir / 'batch_results.json').write_text(json.dumps(results, ensure_ascii=False, indent=2), encoding='utf-8')
    with (output_dir / 'processing_report.csv').open('w', newline='', encoding='utf-8-sig') as handle:
        writer = csv.DictWriter(handle, fieldnames=['excel_row', 'video_id', 'verdict', 'status', 'error'])
        writer.writeheader()
        writer.writerows(report)
    print('\nSummary:', counts)
    print('JSON:', output_dir / 'batch_results.json')
    print('Report:', output_dir / 'processing_report.csv')
    print('JPEG files: evidence/ (as written by evidence_generator.py)')


if __name__ == '__main__':
    main()
